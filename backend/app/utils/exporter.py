"""论文导出：把 Paper（含 references）转成 docx / pdf 文件字节。
Word 用 python-docx，PDF 用 fpdf2（需系统中文字体）。"""
import os
import re
from io import BytesIO

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")

# Windows 常见中文字体，按优先级取第一个存在的（都是单体 .ttf，fpdf2 能稳定解析）
_CJK_FONT_CANDIDATES = [
    r"C:\Windows\Fonts\simhei.ttf",
    r"C:\Windows\Fonts\simfang.ttf",
    r"C:\Windows\Fonts\simkai.ttf",
]


def _add_runs(paragraph, text):
    """把 **加粗** 文本拆成 normal / bold 两种 run 加入段落"""
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        run = paragraph.add_run(m.group(1))
        run.bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _plain(text):
    """去掉 markdown 加粗符号，作为纯文本"""
    return _BOLD_RE.sub(r"\1", text or "")


def _ref_lines(references):
    """参考文献 → ['[1] authors. title. journal（year）', ...]"""
    out = []
    for i, r in enumerate(references or [], 1):
        authors = getattr(r, "authors", None)
        title = getattr(r, "title", "") or ""
        journal = getattr(r, "journal", None)
        year = getattr(r, "year", None)
        parts = [p for p in (authors, title, journal) if p]
        line = "[{}] ".format(i) + ". ".join(parts)
        if year:
            line += "（{}）".format(year)
        out.append(line)
    return out


# ---------- Word ----------

def _markdown_to_docx(doc, content):
    for raw in (content or "").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        stripped = line.strip()
        if stripped.startswith("#"):
            level = min(len(stripped) - len(stripped.lstrip("#")), 4)
            text = stripped.lstrip("#").strip()
            doc.add_heading(text, level=level)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, stripped[2:].strip())
        else:
            p = doc.add_paragraph()
            _add_runs(p, stripped)


def build_docx(paper) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(paper.title or "", level=0)

    if getattr(paper, "abstract", None):
        p = doc.add_paragraph()
        p.add_run("摘要：").bold = True
        p.add_run(paper.abstract)
    if getattr(paper, "keywords", None):
        p = doc.add_paragraph()
        p.add_run("关键词：").bold = True
        p.add_run(paper.keywords)
    if getattr(paper, "topic", None):
        p = doc.add_paragraph()
        p.add_run("选题方向：").bold = True
        p.add_run(paper.topic)

    if getattr(paper, "content", None):
        _markdown_to_docx(doc, paper.content)

    refs = getattr(paper, "references", None) or []
    if refs:
        doc.add_heading("参考文献", level=1)
        for line in _ref_lines(refs):
            doc.add_paragraph(line)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- PDF ----------

def _find_cjk_font():
    for p in _CJK_FONT_CANDIDATES:
        if os.path.exists(p):
            return p
    return None


def _wrap_cjk(pdf, text):
    """按当前字体实际宽度把文本逐字符切成不超过可用宽度的行。
    绕开 fpdf2 wrapmode='CHAR' 在短行上的死循环 bug，中英混排也安全。"""
    available = pdf.w - pdf.l_margin - pdf.r_margin
    lines = []
    for para in text.split("\n"):
        cur = ""
        for ch in para:
            nxt = cur + ch
            if cur and pdf.get_string_width(nxt) > available:
                lines.append(cur)
                cur = ch
            else:
                cur = nxt
        lines.append(cur)
    return lines


def _print_cjk(pdf, text, size, h, align="L"):
    """以 size 字号打印一段（可能含换行）CJK 文本，自动按宽度切行。"""
    pdf.set_font("cjk", "", size)
    for line in _wrap_cjk(pdf, text):
        pdf.cell(0, h, line, new_x="LMARGIN", new_y="NEXT", align=align)


def _markdown_to_pdf(pdf, content):
    for raw in (content or "").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        stripped = line.strip()
        if stripped.startswith("#"):
            level = min(len(stripped) - len(stripped.lstrip("#")), 4)
            text = _plain(stripped.lstrip("#").strip())
            size = {1: 14, 2: 13, 3: 12}.get(level, 12)
            _print_cjk(pdf, text, size, 8)
            pdf.ln(1)
        else:
            if stripped.startswith("- ") or stripped.startswith("* "):
                text = "· " + _plain(stripped[2:].strip())
            else:
                text = _plain(stripped)
            _print_cjk(pdf, text, 11, 7)


def build_pdf(paper) -> bytes:
    from fpdf import FPDF

    font = _find_cjk_font()
    if not font:
        raise RuntimeError("未找到可用的中文字体，无法生成 PDF")

    pdf = FPDF()
    pdf.add_font("cjk", "", font)
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()

    _print_cjk(pdf, paper.title or "", 16, 10, align="C")
    pdf.ln(3)

    if getattr(paper, "abstract", None):
        _print_cjk(pdf, "摘要：" + paper.abstract, 11, 7)
        pdf.ln(2)
    if getattr(paper, "keywords", None):
        _print_cjk(pdf, "关键词：" + paper.keywords, 11, 7)
        pdf.ln(2)
    if getattr(paper, "topic", None):
        _print_cjk(pdf, "选题方向：" + paper.topic, 11, 7)
        pdf.ln(2)

    if getattr(paper, "content", None):
        _markdown_to_pdf(pdf, paper.content)

    refs = getattr(paper, "references", None) or []
    if refs:
        pdf.ln(4)
        _print_cjk(pdf, "参考文献", 13, 8)
        pdf.ln(2)
        for line in _ref_lines(refs):
            _print_cjk(pdf, _plain(line), 11, 7)

    return bytes(pdf.output())